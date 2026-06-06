"""Document loaders for PDF, DOCX, TXT, Markdown."""

from __future__ import annotations

import shutil
import subprocess
import tempfile
from collections.abc import Iterable
from dataclasses import dataclass, field
from pathlib import Path
from typing import TYPE_CHECKING, Any

from app.common.logging import get_logger

if TYPE_CHECKING:
    from app.config import IngestionSection

log = get_logger(__name__)

SUPPORTED_EXTS = {"pdf", "docx", "txt", "md", "markdown"}
_OCR_COMMAND_TIMEOUT_SECONDS = 120


@dataclass(slots=True)
class LoadedPage:
    """로더가 반환하는 문서의 한 단위 (페이지 또는 청크 이전의 텍스트)."""

    source: str
    doc_type: str
    page: int | None
    location: str
    text: str
    metadata: dict[str, object] = field(default_factory=dict)


def detect_doc_type(path: str | Path) -> str:
    suffix = Path(path).suffix.lower().lstrip(".")
    if suffix == "markdown":
        return "md"
    return suffix


def is_supported(path: str | Path) -> bool:
    return detect_doc_type(path) in SUPPORTED_EXTS


def _default_ingestion_settings() -> IngestionSection:
    from app.config import IngestionSection

    return IngestionSection()


def _truncate_text(text: str, limit: int) -> tuple[str, bool]:
    if limit <= 0 or len(text) <= limit:
        return text, False
    return text[:limit], True


def _image_dimensions(image_info: tuple[Any, ...]) -> tuple[int | None, int | None]:
    if len(image_info) < 4:
        return None, None
    width = image_info[2] if isinstance(image_info[2], int) else None
    height = image_info[3] if isinstance(image_info[3], int) else None
    return width, height


def _image_meets_area_threshold(width: int | None, height: int | None, threshold: int) -> bool:
    if threshold <= 0:
        return True
    if width is None or height is None:
        return True
    return width * height >= threshold


def _count_new_significant_images(
    image_infos: list[tuple[Any, ...]],
    threshold: int,
    seen_xrefs: set[int],
) -> int:
    count = 0
    for image_info in image_infos:
        xref = image_info[0] if image_info else None
        if not isinstance(xref, int) or xref in seen_xrefs:
            continue
        if _image_meets_area_threshold(*_image_dimensions(image_info), threshold):
            seen_xrefs.add(xref)
            count += 1
    return count


def _render_page_png(fitz: Any, page_obj: Any, page_number: int, dpi: int, root: Path) -> Path:
    target = root / f"page-{page_number:03d}.png"
    matrix = fitz.Matrix(dpi / 72, dpi / 72)
    pix = page_obj.get_pixmap(matrix=matrix, alpha=False)
    pix.save(str(target))
    return target


def _run_tesseract(render_path: Path, language: str) -> tuple[str, str]:
    if not shutil.which("tesseract"):
        return "", "tesseract is unavailable"
    try:
        result = subprocess.run(
            ["tesseract", str(render_path), "stdout", "-l", language],
            capture_output=True,
            text=True,
            timeout=_OCR_COMMAND_TIMEOUT_SECONDS,
        )
    except subprocess.TimeoutExpired:
        return "", f"tesseract timed out after {_OCR_COMMAND_TIMEOUT_SECONDS} seconds"
    if result.returncode != 0:
        return "", result.stderr.strip() or "tesseract failed"
    return result.stdout, ""


def _combine_pdf_text(text: str, ocr_text: str) -> str:
    text = text.strip()
    ocr_text = ocr_text.strip()
    if text and ocr_text:
        return f"{text}\n\n[OCR]\n{ocr_text}"
    return text or ocr_text


def _load_pdf_pypdf(path: Path) -> list[LoadedPage]:
    try:
        from pypdf import PdfReader
    except ImportError as e:  # pragma: no cover
        raise RuntimeError("pypdf가 설치되어 있지 않습니다. 'pip install pypdf'") from e

    reader = PdfReader(str(path))
    pages: list[LoadedPage] = []
    for idx, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception as e:
            log.warning("PDF 페이지 추출 실패 (%s p.%d): %s", path.name, idx, e)
            text = ""
        text = text.strip()
        if not text:
            continue
        pages.append(
            LoadedPage(
                source=path.name,
                doc_type="pdf",
                page=idx,
                location=f"page={idx}",
                text=text,
                metadata={
                    "parser_backend": "pypdf",
                    "text_chars": len(text),
                    "ocr_attempted": False,
                    "ocr_used": False,
                    "needs_visual_review": False,
                },
            )
        )
    return pages


def _load_pdf_pdf_monster(path: Path, ingestion: IngestionSection) -> list[LoadedPage]:
    try:
        import fitz  # type: ignore
    except ImportError as e:  # pragma: no cover
        raise RuntimeError("PyMuPDF가 설치되어 있지 않습니다. 'pip install PyMuPDF'") from e

    doc = fitz.open(str(path))
    try:
        if getattr(doc, "needs_pass", False):
            raise RuntimeError("PDF is encrypted and requires a password")
        page_count_value = getattr(doc, "page_count", None)
        page_count = int(page_count_value if page_count_value is not None else len(doc))
        pages: list[LoadedPage] = []
        seen_visual_review_xrefs: set[int] = set()

        with tempfile.TemporaryDirectory(prefix="omn-pdf-") as tmp_dir:
            render_root = Path(tmp_dir)
            for page_index in range(page_count):
                page_number = page_index + 1
                page_warnings: list[str] = []
                page_obj = doc.load_page(page_index)

                try:
                    raw_text = page_obj.get_text("text", sort=True) or ""
                except Exception as exc:
                    raw_text = ""
                    page_warnings.append(f"text extraction failed: {exc}")

                raw_text = raw_text.strip()
                text_chars = len(raw_text)
                text, text_truncated = _truncate_text(raw_text, ingestion.pdf_max_page_text_chars)

                try:
                    page_images = page_obj.get_images(full=True)
                except Exception as exc:
                    page_images = []
                    page_warnings.append(f"image inspection failed: {exc}")

                new_significant_image_count = _count_new_significant_images(
                    page_images,
                    ingestion.pdf_visual_review_image_area,
                    seen_visual_review_xrefs,
                )
                should_ocr = ingestion.pdf_ocr == "always" or (
                    ingestion.pdf_ocr == "auto" and text_chars < ingestion.pdf_ocr_threshold
                )

                raw_ocr_text = ""
                ocr_text = ""
                ocr_text_truncated = False
                if should_ocr:
                    try:
                        render_path = _render_page_png(
                            fitz,
                            page_obj,
                            page_number,
                            ingestion.pdf_render_dpi,
                            render_root,
                        )
                        raw_ocr_text, warning = _run_tesseract(render_path, ingestion.pdf_ocr_lang)
                    except Exception as exc:
                        raw_ocr_text = ""
                        warning = f"OCR render failed: {exc}"
                    if warning:
                        page_warnings.append(warning)
                    ocr_text, ocr_text_truncated = _truncate_text(
                        raw_ocr_text.strip(),
                        ingestion.pdf_max_page_text_chars,
                    )

                visual_review_reasons: list[str] = []
                if text_chars < ingestion.pdf_ocr_threshold:
                    visual_review_reasons.append("sparse_text")
                if text_truncated or ocr_text_truncated:
                    visual_review_reasons.append("text_truncated")
                if should_ocr:
                    visual_review_reasons.append("ocr_attempted")
                if new_significant_image_count:
                    visual_review_reasons.append(
                        f"new_significant_embedded_images:{new_significant_image_count}"
                    )

                combined_text = _combine_pdf_text(text, ocr_text)
                if not combined_text:
                    continue

                pages.append(
                    LoadedPage(
                        source=path.name,
                        doc_type="pdf",
                        page=page_number,
                        location=f"page={page_number}",
                        text=combined_text,
                        metadata={
                            "parser_backend": "pdf_monster",
                            "text_chars": text_chars,
                            "text_truncated": text_truncated,
                            "embedded_image_count": len(page_images),
                            "ocr_attempted": should_ocr,
                            "ocr_used": bool(ocr_text.strip()),
                            "ocr_text_chars": len(raw_ocr_text.strip()),
                            "needs_visual_review": bool(visual_review_reasons),
                            "visual_review_reasons": ",".join(visual_review_reasons),
                            "parse_warnings": " | ".join(page_warnings),
                        },
                    )
                )
        return pages
    finally:
        doc.close()


def _load_pdf(path: Path, ingestion: IngestionSection | None = None) -> list[LoadedPage]:
    settings = ingestion or _default_ingestion_settings()
    if settings.pdf_backend == "pypdf":
        return _load_pdf_pypdf(path)

    try:
        pages = _load_pdf_pdf_monster(path, settings)
    except Exception as exc:
        if settings.pdf_backend == "pdf_monster":
            raise
        log.warning("PDF Monster 파서 실패, pypdf로 폴백 (%s): %s", path.name, exc)
        return _load_pdf_pypdf(path)

    if pages or settings.pdf_backend == "pdf_monster":
        return pages

    log.warning("PDF Monster 파서가 텍스트를 추출하지 못해 pypdf로 폴백합니다 (%s)", path.name)
    return _load_pdf_pypdf(path)


def _load_docx(path: Path) -> list[LoadedPage]:
    try:
        from docx import Document
    except ImportError as e:  # pragma: no cover
        raise RuntimeError("python-docx가 설치되어 있지 않습니다.") from e

    doc = Document(str(path))
    paragraphs: list[str] = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            paragraphs.append(t)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                paragraphs.append(" | ".join(cells))
    body = "\n".join(paragraphs).strip()
    if not body:
        return []
    return [
        LoadedPage(
            source=path.name,
            doc_type="docx",
            page=None,
            location="section=body",
            text=body,
        )
    ]


def _load_text(path: Path, doc_type: str) -> list[LoadedPage]:
    for encoding in ("utf-8", "utf-8-sig", "cp949", "euc-kr"):
        try:
            text = path.read_text(encoding=encoding)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover
        text = path.read_text(encoding="utf-8", errors="replace")
    text = text.strip()
    if not text:
        return []
    return [
        LoadedPage(
            source=path.name,
            doc_type=doc_type,
            page=None,
            location="section=body",
            text=text,
        )
    ]


def load_document(
    path: str | Path,
    relative_path: str | None = None,
    ingestion: IngestionSection | None = None,
) -> list[LoadedPage]:
    """지원 형식의 파일에서 페이지 단위 텍스트를 추출한다.

    :param relative_path: 지정 시 LoadedPage.source 를 이 값으로 덮어쓴다.
                          Vault 모드에서 vault 루트 기준 상대 경로를 전달.
    """
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(f"파일을 찾을 수 없습니다: {p}")
    doc_type = detect_doc_type(p)
    if doc_type not in SUPPORTED_EXTS:
        raise ValueError(f"지원하지 않는 형식: .{doc_type} (파일: {p.name})")

    if doc_type == "pdf":
        pages = _load_pdf(p, ingestion=ingestion)
    elif doc_type == "docx":
        pages = _load_docx(p)
    elif doc_type in ("md", "markdown"):
        pages = _load_text(p, "md")
    else:
        pages = _load_text(p, "txt")

    if relative_path is not None:
        pages = [
            LoadedPage(
                source=relative_path,
                doc_type=pg.doc_type,
                page=pg.page,
                location=pg.location,
                text=pg.text,
                metadata=dict(pg.metadata),
            )
            for pg in pages
        ]
    return pages


def load_documents(
    paths: Iterable[str | Path],
    ingestion: IngestionSection | None = None,
) -> list[LoadedPage]:
    pages: list[LoadedPage] = []
    for p in paths:
        pages.extend(load_document(p, ingestion=ingestion))
    return pages
